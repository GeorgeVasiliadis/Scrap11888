import openpyxl

def importNames(filename, logger):
    names = []

    try:
        workbook = openpyxl.load_workbook(filename)
        sheet = workbook.active

        for i in range(1,sheet.max_row+1):
            names.append(sheet.cell(row=i, column=1).value)
    except:
        logger.log("An error occured while trying to process a file.", "error")
    return names
